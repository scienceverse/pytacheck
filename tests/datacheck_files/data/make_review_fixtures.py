"""Byte-level fixtures for the datacheck_files_review parity cases.

Run from this directory: ``python make_review_fixtures.py``. Each file targets
one R behaviour a reviewer found the port got wrong (see
``tests/datacheck_files/gen_review_cases.py``); make_review_fixtures.R writes
the R-generated ones (haven, readRDS).
"""

from __future__ import annotations

from pathlib import Path

HERE = Path(__file__).resolve().parent / "review"
HERE.mkdir(exist_ok=True)


def w(name: str, data: bytes | str) -> None:
    (HERE / name).write_bytes(data.encode("utf-8") if isinstance(data, str) else data)


# text_peek(): rawToChar() drops trailing NULs; iconv() output with a NUL is NA
w("tp_trailing_nul.csv", b"a,b\n1,2\n\0\0")
w("tp_trailing_nul1.csv", b"a,b\n1,2\n\0")
w("tp_dbl_nul.txt", b"A\0\0\0B\0")
w("tp_all_nul.txt", b"\0\0\0\0")
w("tp_mid_nul.txt", b"x,y\n1,2\n\0\0z")

# readLines() drops a UTF-8 BOM from the first line of every call
w("bom_numeric.csv", b"\xef\xbb\xbf1,2,3\n4,5,6\n7,8,9\n")
w("bom_comment_cr.txt", b"\xef\xbb\xbf# comment line\rid|x\r1|a\r2|b\r3|c\r")
w("bom_nohdr_blob.csv", b"\xef\xbb\xbf" + b"x" * 10 + b"\n" + b"y" * 5000 + b"\n")

# fread countfields(): an invalid line leaves the position where it was
w(
    "tab_in_quotes.tsv",
    "﻿NA,1\n"
    + "".join(f"{v},00:00:01\n" for v in ["F", "T", "1", "0", "TRUE"] * 6)
    + '"tab\there",00:00:01',
)
# fread: a column name with a NUL is an error (-> read.delim), NULs in values are dropped
w("fread_nul_name.csv", b"tex\0t\nx y\n1i\nTRUE\n")
w("fread_nul_value.csv", b"a,b\n1,x\0y\n2,z\n3,\0w\n")
# fread's first-line field-count assertion (\r-only file, whitespace first line)
w("fread_cr_ws_first.tsv", b'\t\rNA\r\r""\r  x\r')
# integer64 columns carry class "integer64"
w("int64.csv", "id,big,small\n1,12345678901,5\n2,3,6\n3,,7\n")

# read.delim() fallback (fread: "Single column input contains invalid quotes")
w("rt_quotes.csv", 'value\nalpha\na"b"c\n"x""y"\nhe said "hi"\n"unterminated\n')
w("rt_int.csv", 'v\n"1"2\n 3\n4\n')
w("rt_int_trailing_ws.csv", 'v\n"1"2\n3 \n4\n')
w("rt_lgl.csv", 'v\n"T"RUE\nFALSE\nT\n')
w("rt_lgl_lower.csv", 'v\n"T"RUE\nfalse\n')
w("rt_dbl.csv", 'v\n"1".5\n0x1A\n1e5\n-Inf\nNaN\n')
w("rt_nan_upper.csv", 'v\n"N"AN\nNaN\n')
w("rt_cplx.csv", 'v\n"1"i\n2+3i\n')
w("rt_blank_na.csv", 'v\n"x"y\n""\n\nNA\n \n')
w("rt_na_latin1.csv", 'v\n"a"b\nNA\ncafé\n')
w("rt_mb_error.csv", b'v\n"1"2\n\xe9t\xe9\n')
w("rt_dup_rownames.txt", '# comment line\ny\n0\n1,5\n00\n0012\n0\n"unterminated\n00\n')
w("rt_rownames.txt", 'h\n"a"b\na,1\nb,2\nc,3\n')
w("rt_wrap.csv", 'a\n"x"y\n1,2,3\n4\n5,6\n')
w("rt_nul_quote_header.csv", b'"a"\0b a\tb TRUE he said "hi"\n"a,b"\n')
w("rt_nul_value.csv", b'v\n"x"y\nab\0cd\n"q"\0r\n')
w("rt_cr_cr.csv", b'v\r"x"y\r\r\na\r\rb\r')
# fread fails (NUL in the name), read.delim() reads an NA -> the latin1 re-read
w("rt_fallback_latin1.csv", b'v\0x\n"a"b\nNA\ncaf\xc3\xa9\n')
w("rt_nrows.csv", 'v\n"x"y\n' + "".join(f"r{i}\n" for i in range(10)))
# duplicated column names: each keeps its own class (IDate, integer64)
w("dup_names.csv", "d,d,x\n2020-01-01,5000000000,1\n2020-02-01,6000000000,2\n")


# readxl's isDateFormat() ends its scan at any "g" followed by six characters
# (its "General" shortcut), so a Japanese era date format reads as a number
def _era_date_xlsx() -> None:
    import datetime as dt

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["id", "date_jp", "date_iso", "weight"])
    for i, d in enumerate([dt.date(2019, 5, 1), dt.date(2020, 1, 15)]):
        ws.append([i + 1, d, d, 1.5 + i])
        ws.cell(i + 2, 2).number_format = '[$-411]ggge"年"m"月"d"日"'
        ws.cell(i + 2, 3).number_format = "yyyy-mm-dd"
        ws.cell(i + 2, 4).number_format = '0.0 "kg"'
    wb.properties.created = dt.datetime(2026, 1, 1)
    wb.save(HERE / "era_date.xlsx")


_era_date_xlsx()
