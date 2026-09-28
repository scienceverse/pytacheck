"""Regenerate the byte-level fixtures for the datacheck_files parity cases.

Run from this directory: ``python make_fixtures.py`` (openpyxl and
pyreadstat are needed). make_fixtures.R writes the R-generated files; the
readxl example workbooks (MIT) and haven's iris.sas7bdat are copied as is.
"""

from __future__ import annotations

import datetime as dt
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent


def w(name: str, data: bytes | str) -> None:
    (HERE / name).write_bytes(data.encode("utf-8") if isinstance(data, str) else data)


# -- delimited text ------------------------------------------------------------
w("semicolon.csv", "id;score;grp\n1;1,5;a\n2;2,25;b\n3;3;a\n")
w("tabbed.tsv", "id\tval\tlabel\n1\t10\tx y\n2\t20\t\n3\t\tz\n")
w("quoted.csv", 'id,text,n\n1,"a, b",2\n2,"he said ""hi""",3\n3,"multi\nline",4\n4,"NA",\n')
w("latin1_row3.csv", b"id,sentence\n1,ok\n2,ok\n3,She\xead like\n")
w("latin1_header.csv", b"id,caf\xe9\n1,caf\xe9 study\n2,plain\n")
w("bom.csv", b"\xef\xbb\xbfa,b\n1,x\n2,y\n")
w("crlf.csv", "a,b,c\r\n1,2.5,TRUE\r\n2,3.5,FALSE\r\n")
w("headerless.csv", "1,2,3\n4,5,6\n7,8,9\n")
w("headerless_na.csv", "1,NA,3\n,5,Inf\n7,8,9\n")
w("single_col.csv", "score\n" + "\n".join(str(v) for v in (0.5, -1.25, 3, 2.75, 0)) + "\n")
w("blank_lines.csv", "a,b\n1,2\n3,4\n\n5,6\n")
w("trailing_blank.csv", "a,b\n1,2\n3,4\n\n\n")
w("footer.csv", "a,b,c\n1,2,3\n4,5,6\nTotal rows: 2\n")
w("ragged.csv", "a,b\n1,2\n3,4,5\n6,7\n")
w("title_line.csv", "My data export\na,b,c\n1,2,3\n4,5,6\n")
w("comments.csv", "# exported by tool\n# version 2\nx,y\n1,2\n3,4\n")
w(
    "dates.csv",
    "id,day,stamp\n1,2020-01-01,2020-01-01 10:00:00\n2,2020-02-29,2020-01-02T11:30:00Z\n",
)
w("types.csv", "i,d,l,s,lead\n1,1.5,TRUE,x,007\n2,,FALSE,,010\n3,1e5,NA,NA,00\n")
w("big_int.csv", "id,big\n1,12345678901\n2,\n3,-5\n")
w("na_strings.csv", "a,b,c\nNA,NA,\n,x,\n")
w("empty.csv", b"")
w("whitespace.csv", "   \n  \n")
w("header_only.csv", "a,b,c\n")
w("cr_only.csv", b"a,b\r1,2\r3,4\r")
w("cr_blank_start.csv", b"\ra,b\r1,2\r3,4\r")
w("pipe.txt", "x|y|z\n1|2|3\n4|5|6\n")
w("table.txt", "subject\ttrial\trt\n1\t1\t500\n1\t2\t612\n")
w("prose.txt", "These are notes about the study.\nNothing, really, to see here.\nJust prose.\n")
w("one_line.txt", "a single line, with a comma\n")
w(
    "eprime_utf16.txt",
    (
        "﻿*** Header Start ***\r\nVersionPersist: 1\r\nLevelName: Session\r\n"
        "Experiment: naming\r\nSubject: 1\r\n*** Header End ***\r\n"
    ).encode("utf-16-le"),
)
w("utf16_table.txt", "﻿Subject\tTrial\tRT\n1\t1\t500\n1\t2\t600\n".encode("utf-16-le"))
w("eprime_utf8.txt", "Experiment: naming\nSubject: 1\nLevelName: Trial\n")
w("latin1_notes.txt", b"Notes about caf\xe9 study\nsecond line\n")
blob = '"{' + ",".join(f'""k{i}"":{i}' for i in range(1, 5001)) + '}"'
w("blob.csv", "studyRunData\n" + blob + "\n")
cell = '"[' + ",".join(["1.0"] * 5000) + ']"'
w("bigcells.csv", "id,label,arr\n" + "\n".join(f"{i},r{i},{cell}" for i in range(1, 6)) + "\n")
q = lambda *xs: ",".join(f'"{x}"' for x in xs)  # noqa: E731
qual = [
    q(
        "StartDate",
        "EndDate",
        "Status",
        "Progress",
        "Duration (in seconds)",
        "Finished",
        "ResponseId",
    ),
    q(
        "Start Date",
        "End Date",
        "Response Type",
        "Progress",
        "Duration (in seconds)",
        "Finished",
        "Response ID",
    ),
    q(
        '{"ImportId":"startDate"}',
        '{"ImportId":"endDate"}',
        '{"ImportId":"status"}',
        '{"ImportId":"progress"}',
        '{"ImportId":"duration"}',
        '{"ImportId":"finished"}',
        '{"ImportId":"_recordId"}',
    ),
]
qual += [
    q(f"2021-05-{i:02d} 10:00:00", f"2021-05-{i:02d} 10:05:00", 0, 100, 300, 1, f"R_{i:08d}")
    for i in range(1, 7)
]
w(
    "qualtrics.csv",
    "\n".join(qual).replace('{"ImportId":', '{""ImportId"":').replace('"}"', '""}"') + "\n",
)
w("manifest.csv", "type,file\ncode,Study 1.r\ndata,Study 1.csv\ndoc,notes.txt\n")
w(
    "script_refs.R",
    'd <- read.csv("data/raw_scores.csv")\nsaveRDS(d, "processed/clean.rds")\n'
    'x <- readxl::read_excel(path = "Stimuli.xlsx")\nload("workspace.RData")\n',
)

# -- workbooks -----------------------------------------------------------------
import openpyxl


def book(name: str, rows: list[list[object]], start_row: int = 1, start_col: int = 1) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            if v is not None:
                ws.cell(row=start_row + i, column=start_col + j, value=v)
    wb.save(HERE / f"{name}.xlsx")


book(
    "xl_basic",
    [
        ["id", "score", "name", "ok", "when"],
        [1, 2.5, "a", True, dt.datetime(2020, 1, 1)],
        [2, 3, "b", False, dt.datetime(2020, 1, 2, 10, 30)],
        [3, None, "  c  ", None, None],
    ],
)
book(
    "xl_mixed",
    [
        ["a", "b", "c"],
        [1, "x", dt.datetime(2020, 1, 1)],
        [2.5, 3, 5],
        ["t", True, None],
        [None, "", " "],
    ],
)
book("xl_offset", [["x", "y"], [1, 2], [3, 4]], start_row=3, start_col=2)
book("xl_names", [["", "a", "a", None, "b", " b "], [1, 2, 3, 4, 5, 6]])
book("xl_gaps", [["a", None, "c"], [1, None, 3], [None, None, None], [4, None, 6]])
book(
    "xl_nums",
    [
        ["n", "t"],
        [0.1 + 0.2, 1e-10],
        [123456789012345, 1 / 3],
        [1e20, "x"],
        [-0.5, 12345.678901234567],
    ],
)
book("xl_bools", [["a", "b"], [True, 1], [False, "TRUE"], [None, 2]])
book(
    "xl_dates",
    [
        ["d", "dt", "mix"],
        [dt.date(1900, 1, 1), dt.datetime(1999, 12, 31, 23, 59, 59), dt.date(2020, 5, 5)],
        [dt.date(1900, 3, 1), dt.datetime(2000, 1, 1, 0, 0, 0), 10],
    ],
)
book("xl_header_only", [["a", "b"]])
book("xl_numhdr", [[1, 2.5, True], [1, 2, 3]])
book("xl_title", [["Survey export"], [None], ["id", "score"], [1, 5], [2, 7]])
wb = openpyxl.Workbook()
ws = wb.active
ws.append(["a", "b", "c", "d"])
for i in range(1000):
    ws.append([i, i % 2 == 0, dt.datetime(2020, 1, 1) + dt.timedelta(days=i), f"t{i}"])
ws.append(["x", 5, 44000, 3.5])
ws.append([True, "y", "z", True])
ws.append([dt.datetime(2020, 1, 1), 0, True, dt.datetime(2020, 1, 1)])
wb.save(HERE / "xl_coerce.xlsx")
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "first"
ws.append(["only", "first"])
ws.append([1, 2])
ws2 = wb.create_sheet("second")
ws2.append(["s2a", "s2b", "s2c"])
ws2.append(["x", "y", "z"])
wb.save(HERE / "xl_sheets.xlsx")
wb = openpyxl.Workbook()
wb.active.cell(1, 1, "#DIV/0!").data_type = "e"
wb.active.cell(1, 2, "h")
wb.active.cell(2, 1, 1)
wb.active.cell(2, 2, "=1+1")
wb.save(HERE / "xl_error.xlsx")

# -- OpenDocument (hand-written content.xml) -------------------------------------
ODS_HEAD = (
    '<?xml version="1.0" encoding="UTF-8"?>'
    '<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
    'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
    'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" office:version="1.2">'
    '<office:body><office:spreadsheet><table:table table:name="Sheet1">'
)
ODS_TAIL = "</table:table></office:spreadsheet></office:body></office:document-content>"


def c(value: object = None, vtype: str | None = None, text: str | None = None, rep: int = 1) -> str:
    attrs = f' table:number-columns-repeated="{rep}"' if rep > 1 else ""
    if vtype is not None:
        attrs += f' office:value-type="{vtype}"'
        if vtype in ("float", "percentage", "currency"):
            attrs += f' office:value="{value}"'
        elif vtype == "date":
            attrs += f' office:date-value="{value}"'
        elif vtype == "boolean":
            attrs += f' office:boolean-value="{value}"'
    body = "" if text is None else f"<text:p>{text}</text:p>"
    return f"<table:table-cell{attrs}>{body}</table:table-cell>"


def row(*cells: str, rep: int = 1) -> str:
    attrs = f' table:number-rows-repeated="{rep}"' if rep > 1 else ""
    return f"<table:table-row{attrs}>{''.join(cells)}</table:table-row>"


def ods(name: str, rows: list[str]) -> None:
    with zipfile.ZipFile(HERE / name, "w") as zf:
        zf.writestr("mimetype", "application/vnd.oasis.opendocument.spreadsheet")
        zf.writestr("content.xml", ODS_HEAD + "".join(rows) + ODS_TAIL)
        zf.writestr(
            "META-INF/manifest.xml",
            '<?xml version="1.0" encoding="UTF-8"?><manifest:manifest '
            'xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0">'
            '<manifest:file-entry manifest:full-path="/" '
            'manifest:media-type="application/vnd.oasis.opendocument.spreadsheet"/>'
            '<manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>'
            "</manifest:manifest>",
        )


ods(
    "typed.ods",
    [
        row(
            c(text="id"),
            c(text="score"),
            c(text="pct"),
            c(text="ok"),
            c(text="day"),
            c(text="note"),
        ),
        row(
            c(1, "float", "1"),
            c(2.5, "float", "2.50"),
            c(0.5, "percentage", "50%"),
            c("true", "boolean", "TRUE"),
            c("2020-01-01", "date", "2020-01-01"),
            c(text="  hi  "),
        ),
        row(
            c(2, "float", "2"),
            c(),
            c(0.25, "percentage", "25%"),
            c("false", "boolean", "FALSE"),
            c("2020-02-01", "date", "2020-02-01"),
            c(text="NA"),
        ),
        row(c(3, "float", "3"), c(4, "float", "4"), c(), c(), c(), c(text="x"), c(rep=3)),
        row(c(rep=6), rep=2),
    ],
)
ods(
    "offset.ods",
    [
        row(c(rep=2)),
        row(c(), c(text="a"), c(text="b")),
        row(c(), c(1, "float", "1"), c(text="u")),
        row(c(), c(2, "float", "2"), c(text="v")),
    ],
)
ods(
    "title.ods",
    [
        row(c(text="Survey export")),
        row(c(text="id"), c(text="score")),
        row(c(1, "float", "1"), c(5, "float", "5")),
        row(c(2, "float", "2"), c(7, "float", "7")),
    ],
)

# -- SPSS portable -----------------------------------------------------------------
import pandas as pd
import pyreadstat

pyreadstat.write_por(
    pd.DataFrame({"ID": [1.0, 2.0, 3.0], "SCORE": [2.5, None, 4.0], "NAME": ["a", "b", ""]}),
    str(HERE / "portable.por"),
)
