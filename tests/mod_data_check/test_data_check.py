"""Tests of the Data Check module (port of ``inst/modules/data_check.R``).

Ports the module-level tests of metacheck's ``tests/testthat/test-data-checks.R``
(which run ``module_run(test_paper("x"), "data_check", local_path = d,
local_only = TRUE)``). Here each fixture directory is listed as a synthetic
``repo_check`` output (:func:`run_dir`), so the tests do not depend on the
``repo_check`` port; the ``local_path`` variants run the real pipeline when
``repo_check`` is available. The parity cases (``parity/cases/mod_data_check.yaml``)
compare the complete outputs with R.
"""

from __future__ import annotations

import json
import math
import os
import shutil
import zipfile
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
import pytest

import pytacheck as pc
from pytacheck.module import module_run
from pytacheck.modules import _data_check as h
from pytacheck.report import ReportTable
from pytacheck.utils import local_options

from . import helpers as dc

REPOS = dc.REPOS


# -----------------------------------------------------------------------------
# helpers
# -----------------------------------------------------------------------------


def _dir_table(path: Path, pid: str = "p1") -> pd.DataFrame:
    """A repo_check-like listing of every file under *path* (``local_path`` style)."""
    files = sorted(p for p in path.rglob("*") if p.is_file())
    return pd.DataFrame(
        {
            "paper_id": pd.Series([pid] * len(files), dtype="string"),
            "repo_name": pd.Series([path.name] * len(files), dtype="string"),
            "repo_url": pd.Series([str(path)] * len(files), dtype="string"),
            "file_name": pd.Series([f.name for f in files], dtype="string"),
            "file_path": pd.Series([f.relative_to(path).as_posix() for f in files], dtype="string"),
            "file_url": pd.Series([None] * len(files), dtype="string"),
            "file_location": pd.Series([str(f) for f in files], dtype="string"),
            "file_size": pd.Series([float(f.stat().st_size) for f in files], dtype="float64"),
        }
    )


def _repo_output(table: pd.DataFrame, paper: Any = None) -> Any:
    from pytacheck.module import ModuleOutput

    if paper is None:
        paper = pc.test_paper(["x"])
        paper.paper_id = "p1"
    return ModuleOutput(
        module="repo_check",
        title="Repository Check",
        section="general",
        table=table,
        summary_table=pd.DataFrame({"paper_id": pd.Series([paper.paper_id], dtype="string")}),
        paper=paper,
    )


def run_dir(path: Path, careless: bool = True, **kwargs: Any) -> Any:
    """data_check on every file under *path* (the testthat ``local_path`` runs)."""
    with local_options({"metacheck.llm.use": False, "pytacheck.careless": careless}):
        return module_run(_repo_output(_dir_table(path)), "data_check", **kwargs)


def run_local(path: Path, careless: bool = True, **kwargs: Any) -> Any:
    """The testthat call itself: ``module_run(test_paper("x"), "data_check", local_path = d)``."""
    pytest.importorskip("pytacheck.modules.repo_check")
    paper = pc.test_paper(["x"])
    with local_options({"metacheck.llm.use": False, "pytacheck.careless": careless}):
        return module_run(paper, "data_check", local_path=str(path), local_only=True, **kwargs)


def report_text(out: Any) -> str:
    return "\n".join(b for b in out.report if isinstance(b, str))


def _write_csv(df: pd.DataFrame, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(path, index=False)
    return path


def make_careless_survey(
    prefix: str = "panas",
    n_items: int = 10,
    n_ok: int = 50,
    levels: tuple[int, ...] = (2, 3, 4),
    seed: int = 1,
) -> pd.DataFrame:
    """``make_careless_survey()``: varied respondents, a straightliner and an alternator."""
    rng = np.random.default_rng(seed)
    items = rng.choice(levels, size=(n_ok, n_items))
    straight = np.full((1, n_items), int(np.median(levels)))
    alternating = np.resize([min(levels), max(levels)], (1, n_items))
    data = np.vstack([items, straight, alternating])
    df = pd.DataFrame(data, columns=[f"{prefix}_{i}" for i in range(1, n_items + 1)])
    df.insert(0, "participant_id", range(1, n_ok + 3))
    return df


# -----------------------------------------------------------------------------
# column facets (test-data-checks.R)
# -----------------------------------------------------------------------------


def _demo_repo(tmp_path: Path) -> Path:
    rng = np.random.default_rng(3)
    _write_csv(
        pd.DataFrame(
            {
                "id": range(1, 21),
                "age": rng.integers(18, 66, 20),
                "gender": rng.choice(["Male", "Female"], 20),
                "race": rng.choice(["White", "Black", "Asian", "Other"], 20),
                "score": rng.normal(size=20),
            }
        ),
        tmp_path / "demo_repo" / "data" / "study.csv",
    )
    return tmp_path / "demo_repo"


def _check_demographics(mo: Any) -> None:
    for col in ("representation", "measurement_level", "concept", "role"):
        assert col in mo.table.columns
    sem = dict(zip(mo.table["column_name"], mo.table["concept"], strict=True))
    assert sem["age"] == "age"
    assert sem["gender"] == "gender"
    assert sem["race"] == "race"
    assert pd.isna(sem["score"])
    arow = mo.table[mo.table["column_name"] == "age"].iloc[0]
    assert arow["representation"] == "numeric"
    assert arow["measurement_level"] == "ratio"
    assert arow["unit"] == "years"
    irow = mo.table[mo.table["column_name"] == "id"].iloc[0]
    assert irow["role"] == "identifier"


def test_tags_demographic_columns_in_its_table(tmp_path: Path) -> None:
    _check_demographics(run_dir(_demo_repo(tmp_path)))


def test_tags_demographic_columns_local_path(tmp_path: Path) -> None:
    _check_demographics(run_local(_demo_repo(tmp_path)))


def _qualtrics_repo(tmp_path: Path) -> Path:
    def q(*vals: Any) -> str:
        return ",".join(f'"{v}"' for v in vals)

    rng = np.random.default_rng(5)
    lines = [
        q(
            "StartDate",
            "Status",
            "Progress",
            "Duration (in seconds)",
            "Finished",
            "ResponseId",
            "Q1",
        ),
        q(
            "Start Date",
            "Response Type",
            "Progress",
            "Duration (in seconds)",
            "Finished",
            "Response ID",
            "How happy?",
        ),
        q(
            '{""ImportId"":""startDate""}',
            '{""ImportId"":""status""}',
            '{""ImportId"":""progress""}',
            '{""ImportId"":""duration""}',
            '{""ImportId"":""finished""}',
            '{""ImportId"":""_recordId""}',
            '{""ImportId"":""QID1""}',
        ),
    ]
    for i in range(1, 21):
        lines.append(
            q(
                f"2021-05-{i:02d} 10:00:00",
                0,
                100,
                int(rng.integers(200, 901)),
                1,
                f"R_abc{i:05d}",
                int(rng.integers(1, 6)),
            )
        )
    d = tmp_path / "q_tag_repo"
    (d / "data").mkdir(parents=True)
    (d / "data" / "survey.csv").write_text("\n".join(lines) + "\n", encoding="utf-8")
    return d


def _check_qualtrics(mo: Any) -> None:
    sem = dict(zip(mo.table["column_name"], mo.table["concept"], strict=True))
    assert sem["Duration (in seconds)"] == "qualtrics_duration"
    assert sem["Status"] == "qualtrics_status"
    assert sem["Finished"] == "qualtrics_finished"
    q1 = sem["Q1"]
    assert pd.isna(q1) or not str(q1).startswith("qualtrics_")


def test_tags_qualtrics_metadata_columns(tmp_path: Path) -> None:
    mo = run_dir(_qualtrics_repo(tmp_path))
    _check_qualtrics(mo)
    assert mo.table["is_qualtrics"].all()
    assert len(mo.qualtrics) == 1
    assert "#### Qualtrics Survey Metadata" in report_text(mo)


def test_tags_qualtrics_metadata_columns_local_path(tmp_path: Path) -> None:
    _check_qualtrics(run_local(_qualtrics_repo(tmp_path)))


# -----------------------------------------------------------------------------
# .RData sharing recommendation
# -----------------------------------------------------------------------------


def test_recommends_sharing_data_when_rdata_holds_no_data() -> None:
    mo = run_dir(REPOS / "rdata")
    assert "R workspace file" in mo.summary_text
    assert "analysis_workspace.RData" in mo.summary_text


def test_does_not_flag_rdata_holding_a_data_frame(tmp_path: Path) -> None:
    d = tmp_path / "rdata_ok" / "data"
    d.mkdir(parents=True)
    shutil.copy(REPOS / "rdata" / "data" / "clean_data.RData", d)
    mo = run_dir(tmp_path / "rdata_ok")
    assert "R workspace file" not in mo.summary_text
    assert list(mo.table["column_name"]) == ["id", "score"]


# -----------------------------------------------------------------------------
# careless responding
# -----------------------------------------------------------------------------


def test_flags_a_straightliner_but_not_an_alternating_responder(tmp_path: Path) -> None:
    _write_csv(make_careless_survey(), tmp_path / "dc_careless" / "data" / "survey.csv")
    mo = run_dir(tmp_path / "dc_careless")
    assert "careless" in list(mo.keys())
    assert not mo.careless["respondent"].duplicated().any()
    hit = mo.careless[mo.careless["respondent"] == "51"]
    assert len(hit) == 1
    assert hit["max_longstring"].iloc[0] == 10
    assert "same answer 10 times in a row" in hit["threshold"].iloc[0]
    assert "52" not in set(mo.careless["respondent"])
    assert (mo.careless["reasons"] == "straightlining").all()


def test_does_not_screen_scale_blocks_below_the_item_minimum(tmp_path: Path) -> None:
    flat = pd.DataFrame(
        {"participant_id": range(1, 41), "q_1": [3] * 40, "q_2": [3] * 40, "q_3": [3] * 40}
    )
    _write_csv(flat, tmp_path / "dc_shortblock" / "data" / "survey.csv")
    mo = run_dir(tmp_path / "dc_shortblock")
    assert len(mo.careless) == 0


def test_reports_careless_coverage_limits_when_nothing_is_flagged(tmp_path: Path) -> None:
    rng = np.random.default_rng(7)
    items = pd.DataFrame(
        rng.integers(1, 6, size=(40, 10)), columns=[f"panas_{i}" for i in range(1, 11)]
    )
    items["panas_1"] = [1, 5] * 20
    items.insert(0, "participant_id", range(1, 41))
    _write_csv(items, tmp_path / "dc_clean" / "data" / "survey.csv")
    rp = report_text(run_dir(tmp_path / "dc_clean"))
    assert "#### Careless Responding" in rp
    assert "What this does not cover" in rp
    assert "not evidence that a dataset is free of careless responding" in rp


def test_careless_scale_blocks_split_by_variable_name_prefix(tmp_path: Path) -> None:
    s1 = make_careless_survey("panas", n_items=8, n_ok=40, levels=(1, 2, 3, 4, 5), seed=2)
    s2 = make_careless_survey("rse", n_items=6, n_ok=40, levels=(1, 2, 3, 4, 5), seed=3)
    wide = pd.concat([s1, s2.drop(columns="participant_id")], axis=1)
    _write_csv(wide, tmp_path / "dc_two_scales" / "data" / "survey.csv")
    mo = run_dir(tmp_path / "dc_two_scales")
    scales = {s for x in mo.careless["scales"] for s in x.split("; ")}
    assert any(s.startswith("panas") for s in scales)
    assert any(s.startswith("rse") for s in scales)


def test_does_not_run_careless_without_an_id_or_a_scale_block(tmp_path: Path) -> None:
    s = make_careless_survey().drop(columns="participant_id")
    _write_csv(s, tmp_path / "dc_noid" / "data" / "survey.csv")
    assert len(run_dir(tmp_path / "dc_noid").careless) == 0

    rng = np.random.default_rng(11)
    _write_csv(
        pd.DataFrame(
            {"id": range(1, 41), "rt": rng.normal(500, 50, 40), "age": rng.integers(18, 66, 40)}
        ),
        tmp_path / "dc_nonsurvey" / "data" / "d.csv",
    )
    assert len(run_dir(tmp_path / "dc_nonsurvey").careless) == 0


def test_careless_without_the_indices_reports_the_skip(tmp_path: Path) -> None:
    """metacheck without the ``careless`` package: the opportunity is reported."""
    _write_csv(make_careless_survey(), tmp_path / "dc_careless" / "data" / "survey.csv")
    mo = run_dir(tmp_path / "dc_careless", careless=False)
    assert len(mo.careless) == 0
    assert "careless-response checks were skipped" in report_text(mo)
    assert mo.traffic_light == "green"


def test_longstring_and_irv_match_the_careless_package() -> None:
    rows = [[1.0, 2.0, None], [1.0, 3.0, None], [1.0, 3.0, 5.0]]
    cols = [[r[j] for r in rows] for j in range(3)]  # rows of the careless test data
    assert h._longstring(cols) == [3.0, 2.0, 1.0]
    irv = h._irv([[1.0, 1.0, 1.0], [2.0, 3.0, 3.0]])
    assert irv[0] == 0
    assert math.isclose(irv[1], 0.5773502691896258)
    assert math.isnan(h._irv([[1.0, None, None]])[0])


# -----------------------------------------------------------------------------
# spreadsheet formatting (merged from the former spreadsheet_check)
# -----------------------------------------------------------------------------


def make_excel_repo(tmp_path: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill

    d = tmp_path / "xl_fix"
    (d / "data").mkdir(parents=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["id", "grp", "empty", "val"])
    for i, g, v in zip([1, 2, 3, 4], ["a", "b", "a", "b"], [10, 20, 30, 40], strict=True):
        ws.append([i, g, None, v])
    ws.cell(row=2, column=4).fill = PatternFill("solid", fgColor="FFFFCC00")
    ws.cell(row=3, column=4).fill = PatternFill("solid", fgColor="FF00CCFF")
    ws.merge_cells(start_row=7, start_column=1, end_row=7, end_column=2)
    wb.save(d / "data" / "messy.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.append(["id", "score"])
    for i, s in zip([1, 2, 3], [1.1, 2.2, 3.3], strict=True):
        ws.append([i, s])
    wb.save(d / "data" / "clean.xlsx")
    return d


def test_flags_spreadsheet_colour_merges_and_empty_columns(tmp_path: Path) -> None:
    mo = run_dir(make_excel_repo(tmp_path))
    assert mo.traffic_light == "yellow"
    st = mo.summary_table
    assert st["spreadsheet_file_n"].iloc[0] == 2
    assert st["spreadsheet_flagged_file_n"].iloc[0] == 1
    sheet_finds = mo.findings[mo.findings["column"].isna()]
    assert sheet_finds["check"].str.contains("Colour").any()
    assert sheet_finds["check"].str.contains("Merged").any()
    assert sheet_finds["check"].str.contains("Empty or unnamed").any()
    assert any("examined 2 spreadsheet files" in b for b in mo.report if isinstance(b, str))


def test_spreadsheet_checks_are_clean_when_excel_files_are_clean(tmp_path: Path) -> None:
    from openpyxl import Workbook

    d = tmp_path / "xl_clean"
    (d / "data").mkdir(parents=True)
    wb = Workbook()
    ws = wb.active
    ws.append(["id", "score"])
    for i, s in zip([1, 2, 3], [1.1, 2.2, 3.3], strict=True):
        ws.append([i, s])
    wb.save(d / "data" / "clean.xlsx")
    mo = run_dir(d)
    assert mo.summary_table["spreadsheet_file_n"].iloc[0] == 1
    assert mo.summary_table["spreadsheet_flagged_file_n"].iloc[0] == 0
    checks = mo.findings.loc[mo.findings["column"].isna(), "check"]
    assert not checks.str.contains("Colour|Merged|Empty").any()


_ODS_HEADER = (
    '<?xml version="1.0" encoding="UTF-8"?>'
    "<office:document-content"
    ' xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"'
    ' xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0"'
    ' xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"'
    ' xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0"'
    ' xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0"'
    ' office:version="1.2">'
)


def _write_ods(path: Path, content_xml: str) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    manifest = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0"'
        ' manifest:version="1.2">'
        '<manifest:file-entry manifest:full-path="/" '
        'manifest:media-type="application/vnd.oasis.opendocument.spreadsheet"/>'
        '<manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>'
        "</manifest:manifest>"
    )
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("mimetype", "application/vnd.oasis.opendocument.spreadsheet")
        z.writestr("META-INF/manifest.xml", manifest)
        z.writestr("content.xml", content_xml)
    return path


def test_flags_colour_and_merges_in_ods_files(tmp_path: Path) -> None:
    messy = (
        _ODS_HEADER + "<office:automatic-styles>"
        '<style:style style:name="ceRed" style:family="table-cell">'
        '<style:table-cell-properties fo:background-color="#ff0000"/></style:style>'
        '<style:style style:name="ceWhite" style:family="table-cell">'
        '<style:table-cell-properties fo:background-color="#ffffff"/></style:style>'
        "</office:automatic-styles>"
        '<office:body><office:spreadsheet><table:table table:name="Data">'
        "<table:table-row>"
        '<table:table-cell table:number-columns-spanned="2" table:number-rows-spanned="1" '
        'office:value-type="string"><text:p>banner</text:p></table:table-cell>'
        "<table:covered-table-cell/>"
        '<table:table-cell office:value-type="string"><text:p>val</text:p></table:table-cell>'
        "</table:table-row><table:table-row>"
        '<table:table-cell table:style-name="ceRed" office:value-type="float" office:value="1">'
        "<text:p>1</text:p></table:table-cell>"
        '<table:table-cell table:style-name="ceWhite" office:value-type="float" office:value="2">'
        "<text:p>2</text:p></table:table-cell>"
        '<table:table-cell office:value-type="float" office:value="3"><text:p>3</text:p>'
        "</table:table-cell></table:table-row>"
        '<table:table-row table:number-rows-repeated="3">'
        '<table:table-cell table:number-columns-repeated="3"/></table:table-row>'
        "<table:table-row>"
        '<table:table-cell office:value-type="float" office:value="9"><text:p>9</text:p>'
        '</table:table-cell><table:table-cell table:number-columns-repeated="2"/>'
        "</table:table-row>"
        "</table:table></office:spreadsheet></office:body></office:document-content>"
    )
    d = tmp_path / "ods_fix"
    _write_ods(d / "data" / "messy.ods", messy)
    mo = run_dir(d)
    assert mo.traffic_light == "yellow"
    assert mo.summary_table["spreadsheet_file_n"].iloc[0] == 1
    assert mo.summary_table["spreadsheet_flagged_file_n"].iloc[0] == 1
    sheet_finds = mo.findings[mo.findings["column"].isna()]
    assert sheet_finds["check"].str.contains("Colour").any()
    assert sheet_finds["check"].str.contains("Merged").any()
    assert sheet_finds["detail"].str.contains("A1:B1").any()
    assert any("examined 1 spreadsheet file " in b for b in mo.report if isinstance(b, str))
    # the three blank rows are ONE row element: the repeat counter is expanded
    insp = h.dv_ods_inspect(d / "data" / "messy.ods")
    assert insp["sheets"][0]["empty_rows"] == 3
    assert insp["sheets"][0]["color_cells"] == 1


def test_spreadsheet_checks_are_clean_for_a_clean_ods_file(tmp_path: Path) -> None:
    clean = (
        _ODS_HEADER + '<office:body><office:spreadsheet><table:table table:name="Data">'
        "<table:table-row>"
        '<table:table-cell office:value-type="string"><text:p>id</text:p></table:table-cell>'
        '<table:table-cell office:value-type="string"><text:p>score</text:p></table:table-cell>'
        "</table:table-row><table:table-row>"
        '<table:table-cell office:value-type="float" office:value="1"><text:p>1</text:p>'
        '</table:table-cell><table:table-cell office:value-type="float" office:value="1.1">'
        "<text:p>1.1</text:p></table:table-cell></table:table-row><table:table-row>"
        '<table:table-cell office:value-type="float" office:value="2"><text:p>2</text:p>'
        '</table:table-cell><table:table-cell office:value-type="float" office:value="2.2">'
        "<text:p>2.2</text:p></table:table-cell></table:table-row>"
        "</table:table></office:spreadsheet></office:body></office:document-content>"
    )
    d = tmp_path / "ods_clean"
    _write_ods(d / "data" / "clean.ods", clean)
    mo = run_dir(d)
    assert mo.summary_table["spreadsheet_flagged_file_n"].iloc[0] == 0


# -----------------------------------------------------------------------------
# module paths and outputs
# -----------------------------------------------------------------------------


def test_no_files() -> None:
    mo = dc.dc_run("none")
    assert mo.traffic_light == "na"
    assert mo.summary_text == "We found no files to analyse."
    assert len(mo.table) == 0 and len(mo.structure) == 0


def test_no_readable_tabular_data_returns_the_validation_early_exit() -> None:
    mo = dc.dc_run("nodata")
    # metacheck's early exit returns only the validation half: no summary_text,
    # the module's own traffic light and report replaced
    assert mo.traffic_light == "na"
    assert mo.summary_text is None
    assert mo.dv_summary_text == "We found no readable tabular data files to validate."
    assert mo.report == ""
    assert "findings" not in list(mo.keys())


def test_traffic_lights() -> None:
    assert dc.dc_run("basic").traffic_light == "green"
    assert dc.dc_run("yellow").traffic_light == "yellow"
    assert dc.dc_run("flags").traffic_light == "red"
    assert dc.dc_run("nolocal").traffic_light == "yellow"


def test_output_elements_and_structure_columns() -> None:
    mo = dc.dc_run("basic")
    for k in ("findings", "careless", "demographics", "qualtrics", "structure", "previews",
              "gated_repos", "manifest_path", "group_no_evidence", "na_replace"):  # fmt: skip
        assert k in list(mo.keys())
    for col in ("data_type", "doc_role", "data_format", "group", "referenced_by",
                "tabular_usable", "non_tabular_reason"):  # fmt: skip
        assert col in mo.structure.columns
    assert list(mo.previews) == ["study.csv"]
    assert list(mo.table.columns[:6]) == [
        "paper_id", "repo_url", "source_file", "group", "column_name", "representation",
    ]  # fmt: skip
    assert "ambiguous" not in mo.table.columns and "is_numeric" not in mo.table.columns
    assert mo.summary_table.columns.tolist()[:4] == [
        "paper_id", "data_file_n", "column_n", "empty_col_n",
    ]  # fmt: skip


def test_does_not_mutate_the_repo_check_output() -> None:
    prev = dc.dc_prev("flags")
    before = prev.table.copy(deep=True)
    dc.dc_run("flags", prev=prev)
    pd.testing.assert_frame_equal(prev.table, before)


def test_paper_list_summary_table_is_per_paper() -> None:
    mo = dc.dc_run("paperlist")
    st = mo.summary_table.set_index("paper_id")
    assert list(st.index) == ["p1", "p2", "p3"]
    assert st.loc["p1", "data_file_n"] == 1 and st.loc["p1", "column_n"] == 4
    assert st.loc["p2", "column_n"] == 12
    assert st.loc["p3", "data_file_n"] == 0  # na_replace
    assert st.loc["p2", "flagged_n"] > 0 and st.loc["p1", "flagged_n"] == 0


def test_download_argument() -> None:
    with pytest.raises(pc.module.ModuleError, match="should be one of"):
        dc.dc_run("basic", download="bogus")
    a = dc.dc_run("nolocal_urls", download="none")
    b = dc.dc_run("nolocal_urls", download=False)
    assert a.dv_summary_text == b.dv_summary_text


def test_archives_are_expanded() -> None:
    mo = dc.dc_run("archives", peek_zips=True)
    names = mo.structure["file_name"].tolist()
    assert "bundle.zip" not in names and "results.tar.gz" not in names
    assert {"trials.csv", "codebook.txt", "summary.csv", "extra.csv"} <= set(names)
    assert "face01.png" not in names  # materials are left in the archive
    paths = dict(zip(mo.structure["file_name"], mo.structure["file_path"], strict=True))
    assert paths["trials.csv"] == "bundle.zip/inner/trials.csv"
    # without peek_zips a zip is not expanded (tar/gz always are)
    names2 = dc.dc_run("archives").structure["file_name"].tolist()
    assert "bundle.zip" in names2 and "summary.csv" in names2


def test_downloads_file_urls() -> None:
    mo = dc.dc_run("download")
    locs = dict(zip(mo.structure["file_name"], mo.structure["file_location"], strict=True))
    assert locs["study.csv"] == "<cache>/osf.io_abcde/data/study.csv"
    assert pd.isna(locs["analysis.R"])  # code is not fetched under download = "data"
    assert list(mo.table["column_name"]) == ["id", "age", "gender", "score"]


def test_llm_passes_use_the_model() -> None:
    mo = dc.dc_run("llm", llm="basic")
    types = dict(zip(mo.structure["file_name"], mo.structure["data_type"], strict=True))
    assert types["notes.xyz"] == "documentation"
    concept = dict(zip(mo.table["column_name"], mo.table["concept"], strict=True))
    assert concept["q3"] == "reaction_time"
    assert any(
        isinstance(b, str) and b.startswith("LLM model 'mock/dc' reviewed ambiguous cases (1 file")
        for b in mo.report
    )


def test_manifest_is_written(tmp_path: Path) -> None:
    mo = dc.dc_run("basic", manifest=str(tmp_path))
    paths = mo.manifest_path
    assert paths and all(Path(p).exists() for p in paths)
    m = json.loads(Path(paths[0]).read_text(encoding="utf-8"))
    names = {f["file_name"] for f in m["files"]}
    assert "study.csv" in names

    none = dc.dc_run("none", manifest=str(tmp_path / "none"))
    assert none.manifest_path and Path(none.manifest_path[0]).exists()


def test_distribution_figure_without_matplotlib(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(h, "plotting_available", lambda: False)
    mo = dc.dc_run("basic", plot_distributions=True)
    assert "*Install the `matplotlib` package to see the distribution histograms.*" in mo.report
    assert "#### Distributions" not in report_text(dc.dc_run("basic"))


def test_distribution_figure(monkeypatch: pytest.MonkeyPatch) -> None:
    pytest.importorskip("matplotlib")
    mo = dc.dc_run("basic", plot_distributions=True, max_facets=1)
    txt = report_text(mo)
    assert "#### Distributions" in txt
    assert '<img src="data:image/png;base64,' in txt
    assert "Set `max_facets = 2` to plot them all." in txt


def test_distribution_figure_render_failure(monkeypatch: pytest.MonkeyPatch) -> None:
    import builtins

    real_import = builtins.__import__

    def fail(name: str, *args: Any, **kwargs: Any) -> Any:
        if name.startswith("matplotlib"):
            raise ImportError(name)
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", fail)
    specs = [{"file": "a.csv", "col": "x", "values": [1, 2, 3, 4], "lower": None, "upper": None}]
    assert h.data_validate_dist_facets(specs) == "*Distribution figure could not be rendered.*"
    assert h.data_validate_dist_facets([]) is None


def test_report_tables() -> None:
    tables = dc.dc_run_tables("flags")
    issues = [t for t in tables if list(t.columns) == ["File", "Column", "Issues"]]
    assert len(issues) == 1
    assert issues[0]["Issues"].str.contains("<span title='").all()


# -----------------------------------------------------------------------------
# bibr export schema 12.x papers
# -----------------------------------------------------------------------------


def _same(a: Any, b: Any) -> None:
    ra, rb = a.results(), b.results()
    assert ra.keys() == rb.keys()
    for key in ra:
        va, vb = ra[key], rb[key]
        if isinstance(va, pd.DataFrame):
            pd.testing.assert_frame_equal(va, vb)
        elif isinstance(va, list):
            assert len(va) == len(vb), key
            for x, y in zip(va, vb, strict=True):
                if isinstance(x, ReportTable):
                    pd.testing.assert_frame_equal(x.data, y.data)
                else:
                    assert x == y, key
        elif isinstance(va, dict):
            assert va.keys() == vb.keys(), key
        else:
            assert va == vb, key


def test_bibr12_paper_matches_its_legacy_reading(tmp_path: Path) -> None:
    """A native 12.x paper and the same paper written/read as legacy JSON agree."""
    paper12 = pc.read(dc.ROOT / "tests" / "fixtures" / "bibr_v12_full.json")
    legacy_file = pc.paper_write(paper12, "legacy", tmp_path, schema_version=None)
    legacy = pc.read(legacy_file)
    assert pc.paper_id(paper12) == pc.paper_id(legacy) == ["demo"]
    for name in ("bibr12", "groups", "careless_two"):
        a = dc.dc_run(name, paper=paper12, careless=True)
        b = dc.dc_run(name, paper=legacy, careless=True)
        _same(a, b)


def test_bibr12_parity_golden_agrees_with_the_legacy_golden() -> None:
    """R gives the 12.x paper the same data_check result as a test paper."""

    def golden(cid: str) -> dict[str, Any]:
        g = json.loads(
            (dc.ROOT / "parity" / "golden" / "mod_data_check" / f"{cid}.json").read_text(
                encoding="utf-8"
            )
        )
        return dict(zip(g["value"]["names"], g["value"]["v"], strict=True))

    a, b = golden("data_check.bibr12"), golden("data_check.basic")
    for name in ("traffic_light", "summary_text", "report", "findings", "careless"):
        assert a[name] == b[name], name


# -----------------------------------------------------------------------------
# helpers
# -----------------------------------------------------------------------------


def test_repo_tree_rows_orders_folders_first() -> None:
    rows = h.repo_tree_rows(["b.txt", "A/x.csv", "a/y.csv", None, ""])
    assert rows["text"].tolist() == [
        "├── A/", "│   └── x.csv", "├── a/", "│   └── y.csv", "└── b.txt",
    ]  # fmt: skip
    assert rows["leaf_idx"].tolist()[1] == 1


def test_pad_uses_display_width() -> None:
    # formatC(x, width = -(9 + 2)) pads by display width: "📊 data" is 7
    # columns wide (an emoji is two), "❓ unknown" 10
    assert h._pad(["\U0001f4ca data", "❓ unknown"]) == ["\U0001f4ca data    ", "❓ unknown "]


def test_issue_cell_merges_pii_checks() -> None:
    cell = h.dv_issue_cell(
        ["Personal info (values)", "Personal info (column name)"],
        ["Values look like emails.", "Name suggests email."],
    )
    assert cell.count("<span") == 1
    assert (
        "Personally Identifying Information — Values look like emails; Name suggests email" in cell
    )


def test_excel_inspect_handles_missing_files(tmp_path: Path) -> None:
    assert h.dv_excel_inspect(tmp_path / "nope.xlsx") is None
    assert h.dv_ods_inspect(tmp_path / "nope.ods") is None
    bad = tmp_path / "bad.ods"
    bad.write_text("not a zip")
    assert h.dv_ods_inspect(bad) is None


def test_careless_available_option() -> None:
    with local_options({"pytacheck.careless": False}):
        assert h._careless_available() is False
    assert h._careless_available() is True


def test_fixture_listing_is_current() -> None:
    """scenarios.json lists every fixture file (regenerate with make_cases.py)."""
    sc = json.loads((dc.HERE / "scenarios.json").read_text(encoding="utf-8"))
    for repo in sc["basic"]["repos"]:
        listed = set(repo["files"])
        on_disk = {
            p.relative_to(REPOS / repo["dir"]).as_posix()
            for p in (REPOS / repo["dir"]).rglob("*")
            if p.is_file()
        }
        assert listed == on_disk
    assert os.path.isdir(REPOS / "careless")


# -----------------------------------------------------------------------------
# review fixes
# -----------------------------------------------------------------------------


def test_pad_counts_r_format_widths() -> None:
    # R's format() counts a tab as 2 (its escape "\t"), a C0 control as 4, a
    # backslash as 2, a soft hyphen as 1 and U+1F7F0 (newer than R's width
    # table) as 1
    assert h._disp_width("a\tb") == 4
    assert h._disp_width("a\x01b") == 6
    assert h._disp_width("­x") == 2
    assert h._disp_width("x\U0001f7f0") == 2
    assert h._disp_width("a\\b") == 4
    assert h._disp_width("é日") == 3
    assert h._pad(["a\tb", "abcdef"], extra=2) == ["a\tb    ", "abcdef  "]


def test_integer64_columns_print_as_exact_integers() -> None:
    # fread types integers beyond 32 bits as bit64::integer64, whose
    # as.character() never switches to scientific notation
    out = dc.dc_run("review_types")
    tbl = out.table
    big = tbl.loc[(tbl["source_file"] == "types.csv") & (tbl["column_name"] == "big")]
    assert big["sample_values"].tolist() == ["100000 | 200000 | 300000 | 123456789012 | 1"]
    df = pd.DataFrame({"x": pd.Series([100000, None, 400000000000], dtype="Int64")})
    df.attrs["col_attrs"] = {"x": {"class": "integer64"}}
    assert h.col_chr(df, 0) == ["100000", None, "400000000000"]
    df.attrs = {}
    assert h.col_chr(df, 0) == ["1e+05", None, "4e+11"]


def test_careless_respondent_ids_from_an_integer64_column() -> None:
    out = dc.dc_run("review_int64", careless=True)
    assert out["careless"]["respondent"].tolist() == ["100000000004", "100000000020"]


def test_peek_zips_needs_a_single_true() -> None:
    # isTRUE(peek_zips): a truthy non-logical value does not open the zips
    with_true = dc.dc_run("archives", peek_zips=True)
    with_one = dc.dc_run("archives", peek_zips=1)
    names_true = set(with_true["structure"]["file_name"].tolist())
    names_one = set(with_one["structure"]["file_name"].tolist())
    assert "bundle.zip" not in names_true
    assert "bundle.zip" in names_one


def test_repo_tree_rows_scales_linearly() -> None:
    import time

    paths = [f"data/sub{i % 20}/file{i}.csv" for i in range(4000)]
    paths += [f"flat{i}.txt" for i in range(4000)]
    t = time.perf_counter()
    rows = h.repo_tree_rows(paths)
    assert time.perf_counter() - t < 1.5
    assert len(rows) == 8021
    # the 20 folders first (sub0, sub1, sub10, ...), then the flat files
    assert rows["text"].tolist()[:2] == ["├── data/", "│   ├── sub0/"]
    assert rows["text"].tolist()[-1] == "└── flat999.txt"


def test_stack_stats_matches_bind_rows() -> None:
    from pytacheck._r import bind_rows
    from pytacheck.datacheck.columns import data_col_stats
    from pytacheck.modules.data_check import _stack_stats

    cols = [
        pd.Series([1.0, 2.0, None]),
        pd.Series(["a", "b", None], dtype="string"),
        pd.Series([], dtype="float64"),
    ]
    stats = [data_col_stats(c, c) for c in cols]
    fast = _stack_stats(stats)
    ref = bind_rows(stats).reset_index(drop=True)
    pd.testing.assert_frame_equal(fast, ref)
    # an unexpected shape falls back to bind_rows
    odd = [stats[0], pd.DataFrame({"n": [1], "x": ["y"]})]
    assert list(_stack_stats(odd).columns) == list(bind_rows(odd).columns)
